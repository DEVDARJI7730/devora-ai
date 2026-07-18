import io
import csv
import pdfplumber
import docx
import openpyxl
from services.ai_service import stream_gemini_completions

def extract_text_from_file(file_bytes: bytes, filename: str) -> str:
    """
    Extracts readable text contents from PDF, Word, Excel, CSV, or Text files.
    """
    ext = filename.split(".")[-1].lower()
    
    if ext == "txt":
        return file_bytes.decode("utf-8", errors="ignore")
        
    elif ext == "csv":
        text_stream = io.StringIO(file_bytes.decode("utf-8", errors="ignore"))
        reader = csv.reader(text_stream)
        rows = list(reader)
        return "\n".join([", ".join(row) for row in rows[:100]]) # Limit to first 100 rows
        
    elif ext == "pdf":
        text_content = []
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages[:10]: # Limit to first 10 pages to avoid large token overhead
                page_text = page.extract_text()
                if page_text:
                    text_content.append(page_text)
        return "\n\n".join(text_content)
        
    elif ext in ["docx", "doc"]:
        doc_file = docx.Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc_file.paragraphs]
        return "\n".join(paragraphs[:300]) # Limit paragraphs
        
    elif ext in ["xlsx", "xls"]:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        sheet = wb.active
        rows = []
        for r in list(sheet.iter_rows(values_only=True))[:100]: # Limit to 100 rows
            rows.append(", ".join([str(val) if val is not None else "" for val in r]))
        return "\n".join(rows)
        
    else:
        return "Unsupported file type for automated text extraction."

async def summarize_file_content(text: str, filename: str) -> str:
    """
    Summarizes the extracted text content of a file using Gemini (or simulated response).
    """
    prompt = f"Analyze the following document named '{filename}' and write a concise, executive summary detailing its core points, structure, and key details.\n\nDocument Text:\n{text[:4000]}"
    
    # We can collect streaming response from Gemini helper
    summary_text = ""
    async for chunk in stream_gemini_completions(prompt):
        if "data: " in chunk and not "[DONE]" in chunk:
            try:
                # Remove SSE header
                json_str = chunk.replace("data: ", "").strip()
                payload = json.loads(json_str)
                delta = payload.get("choices", [{}])[0].get("delta", {}).get("content", "")
                summary_text += delta
            except:
                pass
                
    if not summary_text:
        summary_text = (
            f"### Document Analysis: {filename}\n"
            f"- **Character Count**: {len(text)} characters extracted.\n"
            f"- **Core Topic**: General document structure recognized.\n"
            f"\n*(Note: Add your Gemini API key in the bottom-left settings menu for full AI document summarization and Q&A).* "
        )
    return summary_text
